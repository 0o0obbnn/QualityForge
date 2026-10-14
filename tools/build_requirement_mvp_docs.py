from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
import re


OUT = Path("deliverables/2026-08-10")
OUT.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5F6B7A"
LIGHT = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
GREEN = "1F3A5F"
GOLD = "7A5A00"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "222222"
TOTAL_DXA = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_font(run, size=None, bold=None, color=BLACK, name="Calibri", italic=None):
    # Named rendering override: use the installed Unicode font for any run that
    # contains CJK. LibreOffice on macOS ignores eastAsia-only fallbacks in some
    # mixed Chinese/Latin runs, which otherwise produces blank glyphs.
    actual_name = "Arial Unicode MS" if re.search(r"[\u3400-\u9fff]", run.text or "") else name
    run.font.name = actual_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), actual_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), actual_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, 9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    tail = paragraph.add_run(" 页")
    set_font(tail, 9, color=MUTED)


def configure_styles(doc, preset):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10 if preset == "standard_business_brief" else 1.25
    vals = {
        "Heading 1": (16, BLUE, 16 if preset == "standard_business_brief" else 18, 8 if preset == "standard_business_brief" else 10),
        "Heading 2": (13, BLUE, 12 if preset == "standard_business_brief" else 14, 6 if preset == "standard_business_brief" else 7),
        "Heading 3": (12, DARK_BLUE, 8 if preset == "standard_business_brief" else 10, 4 if preset == "standard_business_brief" else 5),
    }
    for name, (size, color, before, after) in vals.items():
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Courier New"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.05


def configure_numbering(doc, compact=False):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(existing_abs or [0]) + 1
    num_id = max(existing_num or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    marker = 269 if compact else 360
    text = 540 if compact else 720
    for level, fmt, text_val in [(0, "bullet", "•"), (1, "bullet", "–")]:
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
        num_fmt = OxmlElement("w:numFmt"); num_fmt.set(qn("w:val"), fmt); lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), text_val); lvl.append(lvl_text)
        suff = OxmlElement("w:suff"); suff.set(qn("w:val"), "tab"); lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), str(text + level * 360)); tabs.append(tab); p_pr.append(tabs)
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), str(text + level * 360)); ind.set(qn("w:hanging"), str(text - marker)); p_pr.append(ind)
        spacing = OxmlElement("w:spacing"); spacing.set(qn("w:after"), "80" if compact else "160"); spacing.set(qn("w:line"), "300" if compact else "280"); spacing.set(qn("w:lineRule"), "auto"); p_pr.append(spacing)
        lvl.append(p_pr)
        r_pr = OxmlElement("w:rPr"); fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), "Calibri"); fonts.set(qn("w:hAnsi"), "Calibri"); r_pr.append(fonts); lvl.append(r_pr)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id)); aid = OxmlElement("w:abstractNumId"); aid.set(qn("w:val"), str(abs_id)); num.append(aid); numbering.append(num)
    return num_id


def add_bullet(doc, text, num_id, level=0):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level)); num_pr.append(ilvl)
    num = OxmlElement("w:numId"); num.set(qn("w:val"), str(num_id)); num_pr.append(num)
    p_pr.append(num_pr)
    r = p.add_run(text); set_font(r, 11)
    return p


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), CALLOUT); p._p.get_or_add_pPr().append(shd)
    for i, line in enumerate(text.strip("\n").splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line); set_font(r, 8.5, color=BLACK, name="Courier New")
    return p


def add_table(doc, headers, rows, widths, header_fill=LIGHT, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); set_font(r, font_size, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value)); set_font(r, font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, title, body, kind="info"):
    color = {"info": INK, "risk": RED, "caution": GOLD, "positive": GREEN}[kind]
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [TOTAL_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title); set_font(r, 10.5, bold=True, color=color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body); set_font(r2, 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def new_doc(title, subtitle, doc_type, version, status, preset):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1); sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492); sec.footer_distance = Inches(0.492)
    configure_styles(doc, preset)
    num_id = configure_numbering(doc, compact=(preset == "compact_reference_guide"))
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run(f"QualityForge  |  {doc_type}"); set_font(r, 9, bold=True, color=MUTED)
    add_page_field(sec.footer.paragraphs[0])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title); set_font(r, 23, bold=True, color=INK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    r2 = p2.add_run(subtitle); set_font(r2, 13, color=MUTED)
    add_table(doc, ["文档类型", "版本", "状态", "日期"], [[doc_type, version, status, str(date(2026, 8, 10))]], [1800, 1500, 2460, 3600], LIGHT, 9.5)
    return doc, num_id


def h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(text); set_font(r, 16, bold=True, color=BLUE)
    return p


def h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    r = p.add_run(text); set_font(r, 13, bold=True, color=BLUE)
    return p


def h3(doc, text):
    p = doc.add_paragraph(style="Heading 3")
    r = p.add_run(text); set_font(r, 12, bold=True, color=DARK_BLUE)
    return p


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix); set_font(r, 11, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_prefix):]); set_font(r2, 11)
    else:
        r = p.add_run(text); set_font(r, 11)
    return p


def page_break(doc):
    doc.add_page_break()


def build_prd():
    doc, bullets = new_doc(
        "需求文档质量分析 MVP",
        "产品需求文档（PRD）｜从需求文本到可审计 Testing IR",
        "产品需求文档",
        "v1.0（评审冻结候选）",
        "待产品 / 架构 / 测试联合评审",
        "standard_business_brief",
    )
    add_callout(doc, "产品裁决", "建议进入 MVP 开发，但仅开发“需求质量分析 + 通过门禁后的 Testing IR”。测试方案、场景、用例、自动化执行与知识图谱均不进入首批开发。", "positive")
    h1(doc, "0. 阅读导航")
    add_table(doc, ["章节", "回答的问题"], [
        ("1–3", "为什么做、给谁用、做到什么程度"),
        ("4–6", "范围、主流程和功能要求"),
        ("7–9", "质量门禁、数据与交互"),
        ("10–13", "Agent/Skill 治理、验收、指标与发布"),
        ("14–15", "风险、待决事项和权威源对齐"),
    ], [1800, 7560])

    h1(doc, "1. 背景与问题定义")
    para(doc, "QualityForge 的长期目标是将非结构化需求转化为可执行测试资产。当前最先需要解决的不是“自动生成更多用例”，而是证明输入需求是否足够明确、可测试、可追溯。若缺少这一层，后续 AI 生成的场景、用例与断言只会把原始缺陷放大，并且难以说明结论来自哪里。")
    h2(doc, "1.1 用户当前痛点")
    for item in [
        "需求存在歧义、遗漏、角色冲突或不可度量的非功能要求，通常到设计或执行阶段才暴露。",
        "PRD 与 OpenAPI 等来源相互矛盾时，人工需要逐段比对，结论缺少统一证据格式。",
        "AI 能给出建议，但若没有确定性门禁、来源溯源和人工复核，无法作为下游测试设计的可信输入。",
        "现有平台设计范围较大，若同时建设解析、生成、执行、RAG 和多 Agent 运行时，MVP 的价值验证会被基础设施工作淹没。",
    ]:
        add_bullet(doc, item, bullets)
    h2(doc, "1.2 产品机会")
    para(doc, "把需求评审变成一个可重复、可解释、可修订的结构化过程：系统先产出带证据的六维评审报告，再由确定性 Gate 裁决；仅当需求满足准入条件时，才编译并验证 Testing IR。")

    h1(doc, "2. 产品定位、用户与核心任务")
    add_table(doc, ["角色", "核心任务", "本 MVP 权限"], [
        ("测试工程师", "导入需求、查看问题、发起澄清、确认评审结论", "创建、查看、处置评审；不能绕过 Gate"),
        ("产品经理 / 需求负责人", "回答澄清问题、修订需求文本", "查看证据、提交修订；不能修改系统裁决日志"),
        ("测试负责人", "确认高风险处置、批准进入 IR", "解决 required / forced 评审任务"),
        ("平台管理员", "管理工作空间与账号", "不参与业务结论审批"),
    ], [1700, 3100, 4560])
    add_callout(doc, "角色边界", "本 MVP 不采用“只有一个测试工程师”的单角色假设。至少需要需求提交者、评审处理者和平台管理员三类职责；实际 RBAC 可在试点阶段合并到同一人，但审计记录仍保留角色动作。", "caution")
    h2(doc, "2.1 核心用户故事")
    for story in [
        "作为测试工程师，我希望导入 Markdown/TXT 或粘贴需求正文，快速获得结构化要素与六维质量报告。",
        "作为需求负责人，我希望每个问题都带原文位置、原因、澄清问题和修改建议，以便直接修订。",
        "作为测试负责人，我希望系统在证据缺失、跨源冲突、不可测试验收或不可度量 NFR 时阻断下游编译。",
        "作为测试工程师，我希望在需求通过后获得可追溯的 Testing IR，并知道每一项规则来自哪一段需求。",
    ]:
        add_bullet(doc, story, bullets)

    h1(doc, "3. 目标、成功条件与非目标")
    add_table(doc, ["类别", "定义"], [
        ("产品目标 G1", "把单份需求文档转换成完整的六维质量评审报告。"),
        ("产品目标 G2", "所有发现都能回到可定位的原始证据，推断项显式标记。"),
        ("产品目标 G3", "确定性 Gate 对同一输入给出可重复裁决，且 fail-closed。"),
        ("产品目标 G4", "需求通过后编译并独立验证 Testing IR，作为后续测试设计输入。"),
        ("成功条件", "5–10 份脱敏真实需求完成闭环试点，能统计发现采纳、误报、修订和门禁数据。"),
    ], [2000, 7360])
    h2(doc, "3.1 首批明确不做")
    for item in [
        "测试策略、测试场景、DesignTestCase、断言和自动化脚本生成。",
        "测试执行、Runner、失败归因、自愈、CI/CD 门禁和 Jira 写入。",
        "PDF OCR、复杂 Word 版式、HTML 原型、图片理解、代码差异分析。",
        "RAG、向量数据库、知识图谱、自由规划式多 Agent 编排、分布式队列。",
        "多租户、LDAP/SSO 与生产级密钥管理；首个试点只需单工作空间受控账号。",
    ]:
        add_bullet(doc, item, bullets)

    h1(doc, "4. MVP 范围")
    h2(doc, "4.1 MVP 0.1：需求质量分析")
    add_table(doc, ["范围项", "MVP 定义"], [
        ("主输入", "单份 Markdown、TXT 或粘贴文本；UTF-8；保留标题与段落定位。"),
        ("辅助输入", "可选 OpenAPI 3.x 文件，仅用于跨源一致性比对。"),
        ("结构提取", "章节、功能点、角色、业务实体、业务规则候选、约束、验收标准、NFR。"),
        ("质量分析", "歧义、一致性、可测试性、完整性、NFR 可度量性、角色/权限确定性。"),
        ("发现字段", "严重级别、维度、原文证据、原因、澄清问题、修改建议、推断标记。"),
        ("裁决", "PASS / PASS_WITH_REVIEW / REJECT。"),
        ("输出", "页面报告 + JSON；支持导出 Markdown。"),
    ], [2100, 7260])
    h2(doc, "4.2 MVP 0.2：Testing IR")
    para(doc, "只有质量 Gate 为 PASS，或 PASS_WITH_REVIEW 已由人工完成必要处置后，系统才能进入 IR 流程。IR 包含 requirements、business_rules、entities、state_machines、interface_contracts、constraints 六个子集，并保留来源清单、冲突和未决问题。")
    h2(doc, "4.3 输入限制")
    for item in [
        "单次只分析一个主需求文档；多份 PRD 合并不在首批范围。",
        "单文件建议不超过 2 MB、正文不超过 200,000 字符；超限返回明确错误，不静默截断。",
        "OpenAPI 仅接受 YAML/JSON，大小建议不超过 5 MB；解析失败不影响主文档保存，但阻止跨源一致性得出 PASS。",
        "上传内容不应包含凭证、生产数据和个人敏感信息；日志只存哈希和脱敏摘要。",
    ]:
        add_bullet(doc, item, bullets)

    h1(doc, "5. 端到端业务流程")
    add_code(doc, "导入需求\n  → 解析与结构化\n  → AI 生成六维评审包\n  → requirement-quality-gate 确定性裁决\n      ├─ REJECT：停止，返回修订\n      ├─ PASS_WITH_REVIEW：创建 ReviewTask，人工处置后重检\n      └─ PASS：允许编译 Testing IR\n  → 规则/领域/状态候选抽取\n  → requirement-compiler 审计融合\n  → ir-validator 独立验证\n      ├─ REJECT：保留版本，回到编译\n      ├─ PASS_WITH_REVIEW：附人工评审清单\n      └─ PASS：发布为下游可读取的已验证 IR")
    h2(doc, "5.1 状态约束")
    add_table(doc, ["对象", "状态 / 裁决", "路由规则"], [
        ("RequirementDocument", "Draft → Parsing → Completed / Failed", "以领域模型现行枚举为准；Completed 只表示解析完成，不等于质量通过。"),
        ("QualityReport", "RUNNING → COMPLETED / FAILED", "MVP 新增建议对象；报告含 verdict。"),
        ("Quality verdict", "PASS / PASS_WITH_REVIEW / REJECT", "只有满足人工处置条件的 PASS 或已闭环 PASS_WITH_REVIEW 可进入 IR。"),
        ("TestingIR", "DRAFT → COMPILING → VALIDATED / FAILED", "VALIDATED 仅由 ir-validator 的可验证结果驱动。"),
        ("ReviewTask", "OPEN → RESOLVED / EXPIRED", "人工 resolution 为 approve / reject / modify。"),
    ], [1800, 2700, 4860])

    h1(doc, "6. 功能需求")
    frs = [
        ("FR-01", "导入需求", "支持 Markdown、TXT 与粘贴文本；创建 RequirementDocument，保留原文、文件名、内容哈希和来源定位。", "P0"),
        ("FR-02", "绑定 OpenAPI", "可在导入时附加或后续绑定一个 OpenAPI 3.x 辅助来源，并记录 sourceId 与可用性。", "P0"),
        ("FR-03", "结构化解析", "提取章节、功能点、角色、实体、规则候选、约束、验收标准与 NFR；不完整解析不得伪装 Completed。", "P0"),
        ("FR-04", "六维评审", "每个维度必须显式评估，即使无发现也要保留 assessed=true。", "P0"),
        ("FR-05", "证据与建议", "每条 finding 包含 code、dimension、severity、sourceTrace、evidence、reason、clarificationQuestion、suggestedRevision。", "P0"),
        ("FR-06", "确定性门禁", "运行 requirement-quality-gate；checker 未运行、异常或输入不全一律不能 PASS。", "P0"),
        ("FR-07", "人工评审", "PASS_WITH_REVIEW 或强制事件创建 ReviewTask；保留处理人、决议、备注、时间和输入哈希。", "P0"),
        ("FR-08", "修订与重检", "修订生成新文档版本/分析 attempt，不覆盖原报告和失败证据。", "P0"),
        ("FR-09", "IR 编译", "通过 Gate 后抽取规则、领域、状态机并融合成 TestingIR，保留所有可用来源贡献。", "P0"),
        ("FR-10", "IR 独立验证", "ir-validator 只验证、不修复；验证者读取不可变快照。", "P0"),
        ("FR-11", "查询与导出", "支持查看原文—发现—IR 的链路，导出 JSON 和 Markdown。", "P1"),
        ("FR-12", "审计与观测", "记录 runId、Agent、Skill 版本、输入/输出哈希、裁决、耗时、错误与人工动作。", "P0"),
    ]
    add_table(doc, ["编号", "能力", "要求", "优先级"], frs, [900, 1500, 6060, 900], LIGHT, 8.8)

    h1(doc, "7. 六维质量模型与门禁")
    dims = [
        ("ambiguity", "歧义", "代词指向不明、模糊量词、条件边界不清", "澄清后可唯一解释"),
        ("consistency", "一致性", "PRD 内部或 PRD/OpenAPI 的术语、状态、字段、规则冲突", "冲突保留多版本并由人工裁决"),
        ("testability", "可测试性", "验收没有可观察输入/行为/结果", "可写成 Given/When/Then 且结果可观测"),
        ("completeness", "完整性", "缺角色、前置、异常、数据约束、状态转换", "关键业务路径的必要信息齐备"),
        ("nfr_measurability", "NFR 可度量性", "“快速、稳定、安全”等无指标或测量条件", "指标、统计口径、环境与时间窗明确"),
        ("role_permission", "角色与权限", "谁能看、改、审、撤销不明确或矛盾", "关键动作有唯一责任角色和权限边界"),
    ]
    add_table(doc, ["代码", "维度", "典型问题", "通过条件"], dims, [1550, 1250, 3500, 3060], LIGHT, 8.8)
    h2(doc, "7.1 严重级别")
    add_table(doc, ["级别", "含义", "默认路由"], [
        ("BLOCKER", "缺失证据、结构不完整、跨源硬冲突、关键角色冲突、不可测试验收或不可度量关键 NFR", "REJECT"),
        ("REVIEW", "可在不发明事实的前提下由人工确认的歧义、低置信推断或非关键缺口", "PASS_WITH_REVIEW"),
        ("INFO", "改进建议，不影响当前可测试性和契约一致性", "不单独阻断"),
    ], [1300, 5560, 2500])
    h2(doc, "7.2 裁决规则")
    for item in [
        "REJECT：六维未完整评估、finding 缺来源证据、存在任一 BLOCKER，或 checker 未成功完成。",
        "PASS_WITH_REVIEW：无 BLOCKER，仅剩 REVIEW finding；必须创建人工评审清单。",
        "PASS：六维全部评估且无 finding。",
        "同一输入、同一 Skill 版本、同一配置重复运行，确定性 Gate 裁决必须一致。",
        "系统不得自动补写缺失业务事实，也不得把多个冲突来源静默合并成单一真相。",
    ]:
        add_bullet(doc, item, bullets)

    h1(doc, "8. Testing IR 契约")
    add_table(doc, ["子集", "最小内容", "可信要求"], [
        ("requirements", "需求编号、陈述、类型、优先级、状态", "每项有 sourceTrace；推断显式标记"),
        ("business_rules", "condition、action、source、confidence", "推断规则 confidence ≤ 0.7；无来源拒绝"),
        ("entities", "实体、属性、关系、约束", "不得发明来源中不存在的实体"),
        ("state_machines", "实体、初始状态、状态、转换、触发条件", "初始状态必须显式提供，不可推断"),
        ("interface_contracts", "路径、方法、请求、响应、错误语义", "OpenAPI 缺失时标记 unavailable，不伪造"),
        ("constraints", "业务、数据、时序、合规和 NFR 约束", "保留可验证条件与来源"),
    ], [1750, 3650, 3960])
    add_callout(doc, "发布条件", "TestingIR.status 只有在 ir-validator 对不可变快照给出 PASS，或给出 PASS_WITH_REVIEW 且配置允许带评审发布时，才能进入 VALIDATED。首批建议对 PASS_WITH_REVIEW 仍要求人工确认。", "caution")

    h1(doc, "9. 页面与交互")
    add_table(doc, ["页面", "主要区域", "关键操作"], [
        ("需求导入", "文本粘贴 / 文件选择、OpenAPI 辅助源、输入限制提示", "保存草稿、开始分析"),
        ("分析进度", "阶段、耗时、当前 Skill、失败原因", "取消、重试失败 attempt"),
        ("质量报告", "总裁决、六维概览、finding 列表、原文定位", "筛选、创建澄清、导出"),
        ("人工评审", "待办、证据、建议修改、影响范围", "approve / reject / modify"),
        ("IR 查看器", "六子集树、来源清单、冲突与未决、版本", "验证、导出 JSON"),
    ], [1700, 4260, 3400])
    h2(doc, "9.1 关键交互规则")
    for item in [
        "默认先展示裁决和阻断原因，再展示模型摘要。",
        "点击 finding 必须定位到原文段落；找不到原文定位时该 finding 不能进入正式报告。",
        "Completed、HTTP 200、文件已生成与模型自评均不得展示为“质量通过”。",
        "重检不覆盖历史，用户可比较任意两个 attempt 的 finding 增减。",
    ]:
        add_bullet(doc, item, bullets)

    h1(doc, "10. Agent 与 Skill 运行设计")
    add_table(doc, ["角色", "本 MVP 职责", "允许 Skill"], [
        ("qf-supervisor", "建档、预检、路由、终止和汇总", "plan-orchestrator"),
        ("qf-requirement-ir", "生成质量评审包和 IR 草稿", "requirement-quality-gate、business-rule-extractor、domain-model-extractor、state-machine-builder、requirement-compiler"),
        ("qf-independent-verifier", "对不可变 IR 快照独立验证并路由人工", "ir-validator、hitl-gatekeeper"),
    ], [2200, 3460, 3700])
    para(doc, "权限遵循最小授权：默认无网络、无 Secret、不得写 docs/** 或 Skill 原型；只在单次 run 目录写草稿。生成者不能批准自己的结果，验证器不能修改被验证资产。")

    h1(doc, "11. 验收标准")
    acs = [
        ("AC-01", "Given 用户提交合法 Markdown；When 分析完成；Then 六个维度均有 assessed 记录，报告可在页面和 JSON 中读取。"),
        ("AC-02", "Given 任一 finding；When 用户打开详情；Then 能定位到 sourceId 和 section，展示不超过必要长度的原文证据。"),
        ("AC-03", "Given 关键验收标准不可观察；When Gate 执行；Then verdict=REJECT，IR 编译节点不启动。"),
        ("AC-04", "Given PRD 与 OpenAPI 字段或状态冲突；When 分析；Then 至少保留两方证据，不自动合并，路由人工或 REJECT。"),
        ("AC-05", "Given checker 未执行、超时或异常；When Supervisor 计算裁决；Then 不得返回 PASS，保留错误与输入哈希。"),
        ("AC-06", "Given 相同输入、Skill 版本和配置；When 重复运行 Gate 10 次；Then verdict 与 finding code 集合完全一致。"),
        ("AC-07", "Given PASS_WITH_REVIEW；When 尚有 OPEN ReviewTask；Then Testing IR 不得进入 VALIDATED。"),
        ("AC-08", "Given 修订需求；When 重新分析；Then 创建新 attempt，旧报告、人工决议和哈希仍可查询。"),
        ("AC-09", "Given Gate PASS；When 编译 IR；Then六子集结构存在，所有必需项 sourceTrace 覆盖率为 100%。"),
        ("AC-10", "Given IR 含未解决 conflict 或缺 sourceTrace；When ir-validator 执行；Then REJECT，且验证器不修改输入快照。"),
        ("AC-11", "Given 用户无工作空间读取权限；When 查询文档、报告或 IR；Then 返回 403，响应不泄露资源是否存在。"),
        ("AC-12", "Given 单次流程任一节点 REJECT/ERROR/TIMEOUT；When 路由后续；Then 后续业务节点不启动，证据可审计。"),
    ]
    add_table(doc, ["编号", "Given / When / Then"], acs, [1000, 8360], LIGHT, 9.0)

    h1(doc, "12. 非功能要求")
    add_table(doc, ["类别", "MVP 验收要求"], [
        ("性能", "在约定试点硬件与 20k 字需求样本上，P95 创建请求 < 1 s；异步分析 P95 < 120 s。测量需记录模型、配置和样本大小。"),
        ("可靠性", "任务失败可重试 1 次；无重复写入；相同 Idempotency-Key 返回同一任务。"),
        ("安全", "JWT 鉴权、工作空间隔离、上传类型/大小校验、日志脱敏、默认无外联和 Secret。"),
        ("可审计", "所有节点、裁决、人工动作、版本和输入/输出哈希可查询；不记录隐藏思维链。"),
        ("可访问性", "核心操作可键盘完成；状态不仅依赖颜色；表格和错误提示有文本标签。"),
        ("兼容性", "首批支持最新两个 Chrome 主版本；API 使用 OpenAPI 3.0.x，JSON 字段 camelCase。"),
    ], [1600, 7760])
    add_callout(doc, "指标声明", "性能数值是本 PRD 的试点验收建议，不是现有实测结论。正式冻结前必须用目标模型和脱敏样本跑一次基线；若成本或时延明显偏离，应调整交互为更明确的异步队列，而不是降低质量门禁。", "caution")

    h1(doc, "13. 试点指标与发布计划")
    h2(doc, "13.1 先测量，再设收益目标")
    add_table(doc, ["指标", "定义", "试点采集方式"], [
        ("发现采纳率", "被人工确认并引发需求修订的 finding / 人工已处理 finding", "ReviewTask resolution + 修订差异"),
        ("误报率", "人工判为无效的 finding / 人工已处理 finding", "reject resolution 原因"),
        ("缺陷发现率", "系统发现且人工确认的问题 / Golden Set 标注问题", "5–10 份双人标注样本"),
        ("首轮通过率", "首次分析 verdict=PASS 的文档占比", "QualityReport attempt=1"),
        ("闭环时长", "创建分析到可进入 IR 的时间", "事件时间戳"),
        ("人工修改率", "VALIDATED 前被人工修改的 IR 项 / IR 总项", "版本差异"),
    ], [1700, 4260, 3400])
    h2(doc, "13.2 分阶段")
    add_table(doc, ["阶段", "交付物", "退出条件"], [
        ("Sprint 0：契约冻结", "PRD、接口、QualityReport schema、Finding code、Golden Set 标注规范", "产品/架构/测试联合签字，OpenAPI 无阻断歧义"),
        ("Sprint 1：质量分析", "导入、解析、六维报告、确定性 Gate、历史 attempt", "AC-01 至 AC-08 通过"),
        ("Sprint 2：IR", "三类抽取、编译、独立验证、IR 查看与导出", "AC-09 至 AC-12 通过"),
        ("试点", "5–10 份脱敏真实需求、指标报告、问题清单", "不存在安全/可信链阻断，决定继续或修订"),
    ], [1800, 4200, 3360])

    h1(doc, "14. 风险、依赖与待决事项")
    add_table(doc, ["类型", "事项", "处理建议"], [
        ("Blocker", "QualityReport / QualityFinding 尚未进入术语、领域、API、DB 四源", "Sprint 0 先补正式契约，不在实现中临时造表。"),
        ("Blocker", "现行 OpenAPI 缺少质量报告接口与明确 operationId", "按配套接口文档冻结后更新 32 号权威源。"),
        ("风险", "AI 发现质量受模型与提示词影响", "Golden Set、版本锁定、人工校准；确定性 Gate 不依赖模型自由裁决。"),
        ("风险", "OpenAPI 与 PRD 的来源优先级可能因业务不同", "不预设统一胜者；展示冲突并由人选择适用来源。"),
        ("依赖", "运行目录、Agent allowlist、Skill 版本与评测入口", "复用 55 号 SOP 和机器可读治理配置。"),
        ("待决", "PASS_WITH_REVIEW 是否允许带评审发布 IR", "首批建议不允许；待真实误报数据后再讨论。"),
    ], [1300, 4300, 3760], LIGHT, 8.8)

    h1(doc, "15. 权威源与一致性说明")
    para(doc, "本 PRD 是面向首批开发的可读需求包，不自动取代项目权威源。字段、枚举、术语和权限发生冲突时，仍按下列真相源裁决；本 PRD 提出的新增 QualityReport / QualityFinding 必须经评审后同步进入正式契约。")
    for source in [
        "术语：docs/01-product/03-术语表Glossary.md",
        "实体与状态：docs/01-product/04-核心领域模型.md",
        "AI 自治与阈值：docs/03-governance/21-AI介入点降级矩阵.md",
        "Agent / Skill 运行：docs/03-governance/55-Agent与Skill运行治理SOP.md",
        "API：docs/04-architecture/32-API规范.yaml",
        "DB：docs/04-architecture/33-数据库Schema设计.md",
        "模块背景：docs/02-prd/11-PRD-TestDesign.md",
    ]:
        add_bullet(doc, source, bullets)
    add_callout(doc, "冻结建议", "确认本 PRD 后，不应直接开始全栈编码。下一步是把 QualityReport / QualityFinding 和接口 operationId 同步进 03、04、32、33 四源，并生成可执行的 OpenAPI 契约测试；完成后才满足 Sprint 1 开发入口条件。", "risk")
    path = OUT / "QualityForge-需求文档质量分析MVP-PRD-v1.0.docx"
    doc.save(path)
    return path


def build_api():
    doc, bullets = new_doc(
        "需求文档质量分析 MVP",
        "REST API 接口文档｜现行契约对齐与冻结建议",
        "接口文档",
        "v1.0（评审冻结候选）",
        "待后端 / 前端 / 测试联合评审",
        "compact_reference_guide",
    )
    add_callout(doc, "契约状态", "本文不是对现有 32-API规范.yaml 的替代。标记“现有”的接口可直接复用；标记“增强/新增”的接口必须先回写 OpenAPI、领域模型与 DB 权威源，再进入开发。", "caution")
    h1(doc, "0. 快速索引")
    add_table(doc, ["内容", "位置"], [
        ("全局约定、鉴权、幂等", "第 1–2 章"),
        ("资源与状态", "第 3 章"),
        ("接口总览与详细契约", "第 4–5 章"),
        ("数据模型、错误码、异步处理", "第 6–8 章"),
        ("安全、契约测试、冻结清单", "第 9–11 章"),
    ], [2600, 6760], BLUE_GRAY)

    h1(doc, "1. 文档范围与基线")
    add_table(doc, ["项目", "约定"], [
        ("Base URL", "/api/v1"),
        ("协议", "HTTPS；REST；JSON；文件上传使用 multipart/form-data"),
        ("规范", "OpenAPI 3.0.x；请求/响应字段 camelCase"),
        ("认证", "Authorization: Bearer <JWT>"),
        ("主资源", "RequirementDocument、QualityReport（建议新增）、QualityFinding（建议新增）、TestingIR、BusinessRule、ReviewTask"),
        ("首批范围", "需求导入、解析、六维质量报告、人工处置、IR 编译/查询/验证结果"),
    ], [2100, 7260], BLUE_GRAY)
    h2(doc, "1.1 契约标记")
    for item in [
        "现有：已存在于 docs/04-architecture/32-API规范.yaml；实现前仍需补 operationId 和示例。",
        "增强：保留现有路径，但补充请求、响应、错误、幂等或语义。",
        "新增：当前 OpenAPI 没有对应资源或路径，必须先完成四源同步。",
    ]:
        add_bullet(doc, item, bullets)

    h1(doc, "2. 全局约定")
    h2(doc, "2.1 通用请求头")
    add_table(doc, ["Header", "必需", "说明"], [
        ("Authorization", "是", "Bearer JWT；除健康检查外全局要求。"),
        ("Content-Type", "是", "application/json 或 multipart/form-data。"),
        ("X-Workspace-Id", "建议", "资源路径未显式包含 workspaceId 时用于隔离校验；不能替代服务端归属校验。"),
        ("Idempotency-Key", "异步 POST 必需", "UUID；同用户、同路径、同 key 在 24 h 内返回同一任务。"),
        ("If-Match", "更新建议", "资源版本号，防止覆盖并发修改。"),
    ], [2300, 1200, 5860], BLUE_GRAY)
    h2(doc, "2.2 响应信封")
    para(doc, "现行 ResponseEnvelope 已定义 code、message、timestamp、requestId。本文建议所有非删除响应增加 data；列表响应增加 page 元数据。")
    add_code(doc, '''{
  "code": "SUCCESS",
  "message": "操作成功",
  "timestamp": 1786329600000,
  "requestId": "req-544d9e96-dbbc-4c98-aae1-b6cbd3ae96dd",
  "data": {}
}''')
    h2(doc, "2.3 HTTP 语义")
    add_table(doc, ["场景", "状态码"], [
        ("创建同步资源", "201 Created"),
        ("受理异步任务", "202 Accepted"),
        ("查询 / 更新成功", "200 OK"),
        ("删除成功", "204 No Content，或保留现行 200 envelope；冻结时二选一，不混用"),
        ("输入错误 / 状态冲突", "400 / 409 / 422"),
        ("未认证 / 无权限 / 不存在", "401 / 403 / 404"),
        ("超限 / 频控 / 服务错误", "413 / 429 / 500 / 503"),
    ], [5400, 3960], BLUE_GRAY)

    h1(doc, "3. 资源与状态")
    add_table(doc, ["资源", "状态 / 裁决", "说明"], [
        ("RequirementDocument", "Draft / Parsing / Completed / Failed", "现有权威枚举；Completed 仅表示解析完成。"),
        ("QualityReport", "RUNNING / COMPLETED / FAILED", "建议新增；verdict 独立为 PASS / PASS_WITH_REVIEW / REJECT。"),
        ("QualityFinding", "OPEN / ACCEPTED / DISMISSED / FIXED", "建议新增处置状态；不得与 ReviewTask 状态混用。"),
        ("ReviewTask", "OPEN / RESOLVED / EXPIRED", "现有；resolution=approve / reject / modify。"),
        ("TestingIR", "DRAFT / COMPILING / VALIDATED / FAILED", "现有；compileVerdict=PASS / PASS_WITH_REVIEW / REJECT。"),
    ], [1900, 3000, 4460], BLUE_GRAY, 8.8)
    add_callout(doc, "语义约束", "RequirementDocument.status=Completed 不能被前端翻译成“需求质量通过”；QualityReport.verdict 与 TestingIR.status 是不同层的状态。", "risk")

    h1(doc, "4. 接口总览")
    endpoints = [
        ("TD-001", "POST", "/test-designs/upload", "现有/增强", "上传 Markdown/TXT，可选附 OpenAPI"),
        ("TD-002", "POST", "/test-designs/requirements", "新增", "粘贴文本创建需求"),
        ("TD-003", "GET", "/test-designs/requirements", "现有/增强", "分页查询需求列表"),
        ("TD-004", "GET", "/test-designs/requirements/{id}", "现有", "查询需求与解析状态"),
        ("TD-005", "POST", "/test-designs/{designId}/parse", "现有/增强", "启动解析与质量分析"),
        ("TD-006", "GET", "/test-designs/{designId}/quality-report", "新增", "获取最新或指定 attempt 报告"),
        ("TD-007", "POST", "/test-designs/{designId}/quality-report/recheck", "新增", "修订后重新分析"),
        ("TD-008", "GET", "/reviews", "现有/增强", "查询人工评审任务"),
        ("TD-009", "PUT", "/reviews/{reviewId}/resolve", "现有/增强", "人工解决评审"),
        ("TD-010", "GET", "/test-designs/{designId}/ir", "现有", "获取当前 TestingIR"),
        ("TD-011", "POST", "/test-designs/{designId}/ir/recompile", "现有/增强", "重新编译 IR"),
        ("TD-012", "GET", "/test-designs/{designId}/rules", "现有/增强", "分页获取 BusinessRule"),
    ]
    add_table(doc, ["ID", "方法", "路径", "状态", "用途"], endpoints, [800, 800, 3600, 1300, 2860], BLUE_GRAY, 8.2)
    h2(doc, "4.1 建议 operationId")
    add_code(doc, "TD-001 uploadRequirementDocument\nTD-002 createRequirementFromText\nTD-003 listRequirementDocuments\nTD-004 getRequirementDocument\nTD-005 startRequirementAnalysis\nTD-006 getRequirementQualityReport\nTD-007 recheckRequirementQuality\nTD-008 listReviewTasks\nTD-009 resolveReviewTask\nTD-010 getTestingIr\nTD-011 recompileTestingIr\nTD-012 listBusinessRules")

    h1(doc, "5. 接口详细契约")
    h2(doc, "5.1 TD-001 上传需求文档")
    add_table(doc, ["项", "定义"], [
        ("接口", "POST /test-designs/upload"),
        ("状态", "现有路径；建议增强请求并将成功码从 200 调整为 201。"),
        ("Content-Type", "multipart/form-data"),
        ("幂等", "Idempotency-Key 必需"),
        ("权限", "workspace:requirement:create"),
    ], [1900, 7460], BLUE_GRAY)
    add_table(doc, ["字段", "类型", "必需", "约束"], [
        ("workspaceId", "uuid", "是", "用户必须具有工作空间写权限"),
        ("file", "binary", "是", ".md/.txt；≤2 MB；UTF-8"),
        ("openapiFile", "binary", "否", ".yaml/.yml/.json；≤5 MB；OpenAPI 3.x"),
        ("title", "string", "否", "1–200 字符；未提供时取文件名"),
    ], [1900, 1600, 1000, 4860], BLUE_GRAY)
    add_code(doc, '''{
  "code": "SUCCESS",
  "data": {
    "id": "5d9b7e1e-3f04-4ecb-96ac-405a98616fb0",
    "workspaceId": "1fb8b3a3-75fb-462f-ae8b-9d19a3c90f39",
    "title": "退款需求 v1",
    "fileName": "refund.md",
    "status": "Draft",
    "contentHash": "sha256:...",
    "createdAt": "2026-08-10T10:00:00+08:00"
  }
}''')

    h2(doc, "5.2 TD-002 粘贴文本创建需求")
    add_table(doc, ["项", "定义"], [
        ("接口", "POST /test-designs/requirements"),
        ("状态", "新增；用于非文件输入。"),
        ("Content-Type", "application/json"),
        ("成功", "201 Created"),
    ], [1900, 7460], BLUE_GRAY)
    add_code(doc, '''{
  "workspaceId": "1fb8b3a3-75fb-462f-ae8b-9d19a3c90f39",
  "title": "退款需求 v1",
  "content": "# 退款规则\\n...",
  "sourceType": "PASTED_TEXT"
}''')
    para(doc, "校验：content 不能为空且不超过 200,000 字符；控制字符、二进制内容或无法按 UTF-8 处理时返回 INPUT_ENCODING_INVALID。")

    h2(doc, "5.3 TD-003 / TD-004 查询需求")
    add_table(doc, ["接口", "关键参数", "响应"], [
        ("GET /test-designs/requirements", "workspaceId 必需；page 默认 1；pageSize 默认 20、最大 100；status 可选", "items + page/pageSize/total"),
        ("GET /test-designs/requirements/{id}", "id: uuid", "RequirementDocument + sourceManifest + latestQualitySummary"),
    ], [3300, 3260, 2800], BLUE_GRAY, 8.8)
    add_callout(doc, "兼容性", "现行列表响应 data 直接为数组。建议冻结为 data.items + 分页元数据；若已有客户端，先新增响应字段并提供一版兼容期。", "caution")

    h2(doc, "5.4 TD-005 启动解析与质量分析")
    add_table(doc, ["项", "定义"], [
        ("接口", "POST /test-designs/{designId}/parse"),
        ("成功", "202 Accepted"),
        ("前置", "RequirementDocument.status 为 Draft 或 Failed；无进行中的 attempt。"),
        ("幂等", "Idempotency-Key 必需"),
        ("副作用", "文档进入 Parsing；创建 AnalysisTask 与 QualityReport attempt。"),
    ], [1900, 7460], BLUE_GRAY)
    add_code(doc, '''{
  "analysisMode": "QUALITY_AND_IR_GATE",
  "sourceIds": ["src-prd", "src-openapi"],
  "requestedOutputs": ["QUALITY_REPORT"]
}''')
    add_code(doc, '''{
  "code": "ACCEPTED",
  "data": {
    "taskId": "c5e9abdf-c905-43a6-b0ef-d5e39c562950",
    "reportId": "927928a8-068b-4cc3-83ec-55de5da1b18e",
    "attempt": 1,
    "status": "RUNNING",
    "pollUrl": "/api/v1/test-designs/5d9b.../quality-report?attempt=1"
  }
}''')

    h2(doc, "5.5 TD-006 获取质量报告")
    add_table(doc, ["参数", "类型", "说明"], [
        ("designId", "path uuid", "需求文档 ID"),
        ("attempt", "query integer", "可选；默认最新；从 1 开始"),
        ("include", "query string", "可选：findings,structure,sourceManifest"),
    ], [2100, 2000, 5260], BLUE_GRAY)
    add_code(doc, '''{
  "id": "927928a8-068b-4cc3-83ec-55de5da1b18e",
  "requirementDocumentId": "5d9b7e1e-3f04-4ecb-96ac-405a98616fb0",
  "attempt": 1,
  "status": "COMPLETED",
  "verdict": "PASS_WITH_REVIEW",
  "dimensions": {
    "ambiguity": {"assessed": true, "blockerCount": 0, "reviewCount": 1},
    "consistency": {"assessed": true, "blockerCount": 0, "reviewCount": 0},
    "testability": {"assessed": true, "blockerCount": 0, "reviewCount": 0},
    "completeness": {"assessed": true, "blockerCount": 0, "reviewCount": 0},
    "nfrMeasurability": {"assessed": true, "blockerCount": 0, "reviewCount": 0},
    "rolePermission": {"assessed": true, "blockerCount": 0, "reviewCount": 0}
  },
  "findings": [{
    "code": "AMBIGUOUS_QUANTIFIER",
    "dimension": "ambiguity",
    "severity": "REVIEW",
    "sourceTrace": "src-prd§2.3",
    "evidence": "系统应尽快退款",
    "reason": "缺少完成时限",
    "clarificationQuestion": "退款应在何时完成？",
    "suggestedRevision": "退款申请通过后 5 分钟内返回受理结果",
    "inference": false
  }]
}''')

    h2(doc, "5.6 TD-007 重新分析")
    add_table(doc, ["项", "定义"], [
        ("接口", "POST /test-designs/{designId}/quality-report/recheck"),
        ("请求", "baseAttempt、reason、可选 sourceIds；Idempotency-Key 必需"),
        ("约束", "创建 attempt+1；不得覆盖旧报告；同一文档最多一个 RUNNING attempt"),
        ("成功", "202 Accepted"),
    ], [1900, 7460], BLUE_GRAY)

    h2(doc, "5.7 TD-008 / TD-009 人工评审")
    add_table(doc, ["接口", "契约要点"], [
        ("GET /reviews", "支持 status、assetType、severity、workspaceId、page、pageSize；MVP 的 assetType 至少支持 requirement、rule、contract。"),
        ("PUT /reviews/{reviewId}/resolve", "body: resolution=approve/reject/modify、resolutionNote、expectedVersion；服务端记录 resolvedBy / resolvedAt。"),
    ], [3400, 5960], BLUE_GRAY)
    add_code(doc, '''{
  "resolution": "modify",
  "resolutionNote": "已将“尽快”改为“5 分钟内返回受理结果”",
  "expectedVersion": 3
}''')
    para(doc, "禁止：AI 使用服务账号代替人工调用 resolve；生成者批准自己的输出；未授权用户处理 required / forced 评审。")

    page_break(doc)
    h2(doc, "5.8 TD-010 获取 Testing IR")
    add_table(doc, ["项", "定义"], [
        ("接口", "GET /test-designs/{designId}/ir"),
        ("状态", "现有"),
        ("成功", "200；无 IR 时 404 IR_NOT_FOUND，质量未准入时 409 QUALITY_GATE_NOT_CLEARED"),
        ("响应", "TestingIR；irContent 保持 JSON object，不使用字符串包 JSON"),
    ], [1900, 7460], BLUE_GRAY)
    add_code(doc, '''{
  "id": "70244a74-9222-4362-95b0-44335f74e171",
  "requirementDocumentId": "5d9b7e1e-3f04-4ecb-96ac-405a98616fb0",
  "irVersion": "1.0.0",
  "status": "VALIDATED",
  "compileVerdict": "PASS",
  "irContent": {
    "requirements": [], "businessRules": [], "entities": [],
    "stateMachines": [], "interfaceContracts": [], "constraints": []
  },
  "declaredSources": [{"sourceId": "src-prd", "availability": "AVAILABLE"}],
  "sourceTrace": "src-prd"
}''')

    h2(doc, "5.9 TD-011 重新编译 IR")
    add_table(doc, ["项", "定义"], [
        ("接口", "POST /test-designs/{designId}/ir/recompile"),
        ("成功", "202 Accepted；返回 status=COMPILING 的新 IR 版本或任务引用"),
        ("前置", "质量 Gate 已准入；所有 required ReviewTask 已 RESOLVED"),
        ("幂等", "Idempotency-Key 必需"),
        ("失败", "409 QUALITY_GATE_NOT_CLEARED / OPEN_REVIEW_TASKS / IR_COMPILE_IN_PROGRESS"),
    ], [1900, 7460], BLUE_GRAY)

    h2(doc, "5.10 TD-012 查询业务规则")
    add_table(doc, ["项", "定义"], [
        ("接口", "GET /test-designs/{designId}/rules"),
        ("参数", "page、pageSize、source、minConfidence、inferred"),
        ("响应", "BusinessRule items + 分页元数据"),
        ("可信约束", "sourceTrace 必需；inferred=true 时 confidence 不得大于 0.7"),
    ], [1900, 7460], BLUE_GRAY)

    h1(doc, "6. 数据模型")
    h2(doc, "6.1 RequirementDocument（增强建议）")
    add_table(doc, ["字段", "类型", "必需", "说明"], [
        ("id", "uuid", "是", "资源 ID"),
        ("workspaceId", "uuid", "是", "工作空间"),
        ("title", "string", "是", "1–200 字符"),
        ("fileName", "string", "文件输入时", "原始文件名"),
        ("rawContent", "string", "服务端存储", "默认不在列表响应返回"),
        ("contentHash", "string", "是", "sha256 内容哈希，建议新增"),
        ("parsedElements", "object", "否", "现行 API 定义为 string，应修正为 object"),
        ("status", "enum", "是", "Draft/Parsing/Completed/Failed"),
        ("createdAt", "date-time", "是", "创建时间"),
    ], [2000, 1700, 1700, 3960], BLUE_GRAY, 8.7)
    h2(doc, "6.2 QualityReport（新增建议）")
    add_table(doc, ["字段", "类型", "说明"], [
        ("id", "uuid", "报告 ID"),
        ("requirementDocumentId", "uuid", "源文档"),
        ("attempt", "integer", "从 1 递增；文档内唯一"),
        ("status", "enum", "RUNNING/COMPLETED/FAILED"),
        ("verdict", "enum?", "完成时为 PASS/PASS_WITH_REVIEW/REJECT；运行中为空"),
        ("dimensions", "object", "六维 assessed、计数和摘要"),
        ("findings", "array", "QualityFinding 列表"),
        ("sourceManifest", "array", "sourceId、类型、可用性、hash"),
        ("skillVersion / configHash", "string", "可重复性与审计"),
        ("startedAt / completedAt", "date-time", "耗时统计"),
    ], [2600, 1800, 4960], BLUE_GRAY, 8.7)
    h2(doc, "6.3 QualityFinding（新增建议）")
    add_table(doc, ["字段", "类型", "约束"], [
        ("code", "string", "稳定 finding code"),
        ("dimension", "enum", "六维之一"),
        ("severity", "enum", "BLOCKER/REVIEW/INFO"),
        ("sourceTrace", "string", "必需；<sourceId>§<section>"),
        ("evidence", "string", "必需；仅保留必要原文"),
        ("reason", "string", "必需"),
        ("clarificationQuestion", "string", "REVIEW/BLOCKER 必需"),
        ("suggestedRevision", "string", "可选；不得发明业务事实"),
        ("inference", "boolean", "推断必须 true"),
        ("status", "enum", "OPEN/ACCEPTED/DISMISSED/FIXED"),
    ], [2600, 1800, 4960], BLUE_GRAY, 8.7)
    h2(doc, "6.4 TestingIR / BusinessRule / ReviewTask")
    para(doc, "沿用现有 32 号 API 定义，但需要补 required、nullable、示例和约束。TestingIR.irContent 必须为 object；BusinessRule.sourceTrace 必需；ReviewTask 的 resolvedBy 建议补入正式契约。")

    h1(doc, "7. 错误模型与业务错误码")
    add_code(doc, '''{
  "code": "QUALITY_GATE_NOT_CLEARED",
  "message": "需求质量门禁尚未准入，不能编译 Testing IR",
  "timestamp": 1786329600000,
  "requestId": "req-...",
  "details": {
    "verdict": "REJECT",
    "reportId": "927928a8-068b-4cc3-83ec-55de5da1b18e"
  }
}''')
    errors = [
        ("INPUT_FORMAT_UNSUPPORTED", "415", "文件类型不支持"),
        ("INPUT_TOO_LARGE", "413", "文件或文本超限"),
        ("INPUT_ENCODING_INVALID", "422", "文本编码不可处理"),
        ("OPENAPI_INVALID", "422", "辅助 OpenAPI 解析失败"),
        ("ANALYSIS_IN_PROGRESS", "409", "已有运行中 attempt"),
        ("QUALITY_GATE_REJECTED", "422", "分析完成但质量被拒绝；响应带 reportId"),
        ("QUALITY_GATE_NOT_CLEARED", "409", "未满足 IR 前置条件"),
        ("OPEN_REVIEW_TASKS", "409", "存在未解决 required/forced 评审"),
        ("IR_COMPILE_IN_PROGRESS", "409", "已有 IR 编译任务"),
        ("IR_VALIDATION_FAILED", "422", "IR 确定性验证不通过"),
        ("VERSION_CONFLICT", "409", "If-Match / expectedVersion 冲突"),
        ("SKILL_TIMEOUT", "503", "Skill 超时；不得转为 PASS"),
        ("SKILL_EXECUTION_ERROR", "500", "Skill/checker 异常；fail-closed"),
    ]
    add_table(doc, ["业务码", "HTTP", "含义"], errors, [3500, 1100, 4760], BLUE_GRAY, 8.5)

    h1(doc, "8. 异步、幂等与并发")
    for item in [
        "POST /parse、/quality-report/recheck、/ir/recompile 均为异步接口，立即返回 202 和 task/report/IR 引用。",
        "MVP 使用轮询，不依赖 WebSocket：运行中建议 2 s 起步、指数退避至 10 s；完成后停止。",
        "相同 Idempotency-Key + 相同主体 + 相同路径 + 相同请求体必须返回同一资源；请求体不同返回 409 IDEMPOTENCY_KEY_REUSED。",
        "每份 RequirementDocument 同时最多一个质量分析 attempt 和一个 IR 编译；重试创建新 attempt，不覆盖旧记录。",
        "人工 resolve 使用 expectedVersion 或 If-Match；冲突时返回 409，不做最后写入者覆盖。",
    ]:
        add_bullet(doc, item, bullets)

    h1(doc, "9. 安全与权限")
    add_table(doc, ["控制项", "要求"], [
        ("资源隔离", "服务端从 JWT 与资源归属校验 workspace/tenant；不信任客户端单独传入的 workspaceId。"),
        ("上传安全", "扩展名、MIME、魔数与大小校验；文件名规范化；存储路径不由客户端决定。"),
        ("敏感数据", "不记录原文全文、JWT、Secret 或 OpenAPI 中的凭证；错误详情脱敏。"),
        ("人工审批", "只有具备 review:resolve 的用户可处理 required/forced 工单；记录 resolvedBy。"),
        ("Agent 权限", "Agent 不直接修改正式需求、IR 或 OpenAPI 权威源；默认无网络和 Secret。"),
        ("审计", "记录 requestId、runId、Skill 版本、输入/输出哈希、裁决、耗时和人工动作。"),
    ], [2200, 7160], BLUE_GRAY)

    page_break(doc)
    h1(doc, "10. 契约测试与验收")
    tests = [
        ("CT-01", "上传合法 Markdown 返回 201、RequirementDocument.status=Draft；同 Idempotency-Key 重放资源 ID 不变。"),
        ("CT-02", "不支持格式或超限分别返回 415/413，且不创建 RequirementDocument。"),
        ("CT-03", "启动分析返回 202；轮询从 RUNNING 到 COMPLETED/FAILED，不能凭 HTTP 200 推导 PASS。"),
        ("CT-04", "QualityReport 六维键齐全；缺一维时服务端不得返回 verdict=PASS。"),
        ("CT-05", "每条 BLOCKER/REVIEW finding 均有 sourceTrace、evidence、reason 和 clarificationQuestion。"),
        ("CT-06", "verdict=REJECT 时调用 /ir/recompile 返回 409 QUALITY_GATE_NOT_CLEARED。"),
        ("CT-07", "有 OPEN required ReviewTask 时 IR 不可进入 VALIDATED。"),
        ("CT-08", "resolve 的 expectedVersion 过期返回 409 VERSION_CONFLICT，原记录不变。"),
        ("CT-09", "TestingIR.irContent 返回 JSON object；六子集存在；sourceTrace 必需项覆盖率 100%。"),
        ("CT-10", "同一分析输入和版本重复执行确定性 Gate，verdict 与 finding code 集合一致。"),
        ("CT-11", "无权限用户查询其他工作空间资源返回 403，响应不含标题或原文。"),
        ("CT-12", "Skill 超时/异常返回失败业务码，状态为 FAILED，后续节点未启动。"),
    ]
    add_table(doc, ["编号", "契约验收"], tests, [1000, 8360], BLUE_GRAY, 8.9)

    h1(doc, "11. 现行 OpenAPI 差距与冻结清单")
    gaps = [
        ("GAP-01", "QualityReport / QualityFinding schema 与查询路径缺失", "Blocker", "新增四源契约"),
        ("GAP-02", "现有操作缺 operationId", "Blocker", "按 TD-001～012 补齐并保证唯一"),
        ("GAP-03", "RequirementDocument.parsedElements 在 API 为 string、领域模型为 JSON", "Blocker", "统一为 object"),
        ("GAP-04", "上传接口只接受单 file，无法表达可选 OpenAPI 来源", "P0", "增强 multipart 或单独建 Source 资源"),
        ("GAP-05", "列表响应无统一分页元数据", "P0", "冻结 items/page/pageSize/total"),
        ("GAP-06", "异步任务使用匿名 object，状态与错误不完整", "P0", "新增 AnalysisTask schema"),
        ("GAP-07", "现有 schema 多数字段未声明 required", "P0", "按创建/响应拆分 schema"),
        ("GAP-08", "ReviewTask 未记录 resolvedBy，接口未声明并发控制", "P0", "补字段和 expectedVersion"),
        ("GAP-09", "删除成功语义为 NoContent 名称但返回 200 envelope", "P1", "冻结 204 或 200，统一全局"),
        ("GAP-10", "API version=1.0.0，新增契约未定义兼容策略", "P1", "明确向后兼容窗口和变更日志"),
    ]
    add_table(doc, ["编号", "差距", "级别", "冻结动作"], gaps, [1000, 4800, 1100, 2460], BLUE_GRAY, 8.4)
    h2(doc, "11.1 开发入口条件")
    for item in [
        "PRD 的 QualityReport / QualityFinding 术语、字段、状态和关系完成产品评审。",
        "03 术语、04 领域、32 API、33 DB 四源同步，且全局搜索无旧枚举冲突。",
        "OpenAPI 语法校验、operationId 唯一性、请求/响应示例和负例契约测试通过。",
        "确定性 Gate 与 IR Validator 的版本、输入 schema、fail-closed 规则已锁定。",
        "Golden Set 至少 5 份完成双人标注；性能指标取得基线，不把建议值冒充实测值。",
    ]:
        add_bullet(doc, item, bullets)
    add_callout(doc, "接口冻结裁决", "现有接口可作为骨架，但还不能直接据此完整开发需求质量分析。完成 GAP-01～03 是最低前置；完成 GAP-01～08 后，才建议前后端并行进入 Sprint 1。", "risk")

    h1(doc, "附录 A. 权威来源")
    for source in [
        "docs/04-architecture/32-API规范.yaml（现行 API 权威）",
        "docs/01-product/03-术语表Glossary.md（术语与状态名）",
        "docs/01-product/04-核心领域模型.md（实体与状态）",
        "docs/04-architecture/33-数据库Schema设计.md（物理持久化）",
        "docs/03-governance/21-AI介入点降级矩阵.md（AI 自治与阈值）",
        "docs/03-governance/55-Agent与Skill运行治理SOP.md（Agent / Skill 权限与运行）",
        "prototypes/skills/*/SKILL.md（单 Skill 输入、输出、Must Not 与裁决）",
    ]:
        add_bullet(doc, source, bullets)
    path = OUT / "QualityForge-需求文档质量分析MVP-接口文档-v1.0.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_prd())
    print(build_api())
